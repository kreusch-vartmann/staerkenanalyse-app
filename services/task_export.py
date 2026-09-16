"""
Export einzelner Beobachtungsaufgaben als PDF und DOCX.

Bewusst getrennt von services/report_generator.py: Der Bericht-Export ist
farbig/gebrandet für Beobachtende (Bildschirm + PDF identisch gestaltet),
der Aufgaben-Export hier ist Absicht schwarz auf weiß, da die Blätter den
Teilnehmenden ausgedruckt vorgelegt werden (siehe Anforderung).

Exportiert werden ausschließlich Titel + die beiden Pflicht-Sektionen
"Aufgabe" und "Rahmenbedingungen" (siehe services/task_normalization.py).
Beobachtungsfokus/Facilitator-Hinweise sind NICHT Teil des Exports.
"""

from __future__ import annotations

from html.parser import HTMLParser
from io import BytesIO


class TaskExportError(RuntimeError):
    """Fehler beim Export einer Aufgabe (PDF oder DOCX)."""


# =============================================================================
# PDF-Export (WeasyPrint, schwarz auf weiß)
# =============================================================================

_PDF_CSS = """
    @page { size: A4; margin: 2.5cm 2cm; }
    body {
        font-family: 'DejaVu Sans', Arial, sans-serif;
        color: #000;
        background: #fff;
        line-height: 1.5;
        font-size: 12pt;
    }
    h1 { font-size: 20pt; margin-bottom: 4pt; }
    .meta { font-size: 10pt; color: #333; margin-bottom: 20pt; }
    h2 { font-size: 14pt; margin-top: 20pt; border-bottom: 1px solid #000; padding-bottom: 4pt; }
    p { margin: 6pt 0; }
    ol, ul { margin: 6pt 0 6pt 18pt; padding: 0; }
    li { margin: 3pt 0; }
    strong { font-weight: bold; }
"""


def build_task_pdf_html(task) -> str:
    """Baut das schlichte, druckfertige HTML für den PDF-Export einer Aufgabe."""
    version = task.current_version
    content = version.content if version else ""

    meta_parts = []
    if task.participant_count:
        meta_parts.append(f"{task.participant_count} Teilnehmende")
    if task.duration_minutes:
        meta_parts.append(f"{task.duration_minutes} Minuten")
    meta_line = " · ".join(meta_parts)

    # Der Content enthält bereits <h2>Titel</h2> - wird hier durch eine
    # eigene <h1> ersetzt, damit Titel/Meta einheitlich gestaltet werden
    # und die beiden Sektionen ("Aufgabe"/"Rahmenbedingungen") als <h2>
    # im Dokument klar hervortreten.
    import re

    body = re.sub(r"<h2[^>]*>.*?</h2>", "", content, count=1, flags=re.DOTALL).strip()

    return f"""<!DOCTYPE html>
<html lang="de">
<head>
<meta charset="utf-8">
<title>{task.title}</title>
<style>{_PDF_CSS}</style>
</head>
<body>
    <h1>{task.title}</h1>
    {f'<p class="meta">{meta_line}</p>' if meta_line else ''}
    {body}
</body>
</html>"""


def task_to_pdf_bytes(task) -> bytes:
    """Rendert eine Aufgabe als PDF (bytes). Erfordert WeasyPrint + Pango/Cairo."""
    try:
        from weasyprint import HTML
    except Exception as e:
        raise TaskExportError(
            "WeasyPrint konnte nicht geladen werden. Bitte installieren Sie die "
            f"erforderlichen System-Bibliotheken (Pango, Cairo). Details: {e}"
        ) from e

    html_string = build_task_pdf_html(task)
    return HTML(string=html_string).write_pdf()


# =============================================================================
# DOCX-Export (python-docx)
# =============================================================================


class _TaskHtmlToDocx(HTMLParser):
    """Wandelt die normalisierte Task-HTML (h2/h3/p/ol/ul/li/strong/em) in ein
    python-docx-Dokument um. Bewusst kein generisches HTML-Rendering: Der
    Content ist durch services/task_normalization.py bereits auf diese
    wenigen Tags beschränkt, ein schlanker eigener Parser genügt und
    vermeidet eine zusätzliche Abhängigkeit.
    """

    def __init__(self, document):
        super().__init__()
        self.document = document
        self._list_stack: list[str] = []  # "ol" oder "ul"
        self._bold = False
        self._italic = False
        self._current_paragraph = None

    def _new_paragraph(self, style: str | None = None):
        self._current_paragraph = self.document.add_paragraph(style=style)
        return self._current_paragraph

    def _add_text(self, text: str):
        if not text or not text.strip():
            return
        if self._current_paragraph is None:
            self._new_paragraph()
        run = self._current_paragraph.add_run(text)
        run.bold = self._bold
        run.italic = self._italic

    def handle_starttag(self, tag, attrs):
        if tag == "h2":
            self.document.add_heading("", level=1)
            self._current_paragraph = self.document.paragraphs[-1]
        elif tag == "h3":
            self.document.add_heading("", level=2)
            self._current_paragraph = self.document.paragraphs[-1]
        elif tag == "p":
            self._new_paragraph()
        elif tag == "ol":
            self._list_stack.append("ol")
        elif tag == "ul":
            self._list_stack.append("ul")
        elif tag == "li":
            list_type = self._list_stack[-1] if self._list_stack else "ul"
            style = "List Number" if list_type == "ol" else "List Bullet"
            self._new_paragraph(style=style)
        elif tag == "strong":
            self._bold = True
        elif tag == "em":
            self._italic = True
        elif tag == "br":
            if self._current_paragraph is not None:
                self._current_paragraph.add_run().add_break()

    def handle_endtag(self, tag):
        if tag in ("ol", "ul") and self._list_stack:
            self._list_stack.pop()
        elif tag == "strong":
            self._bold = False
        elif tag == "em":
            self._italic = False
        elif tag in ("p", "li", "h2", "h3"):
            self._current_paragraph = None

    def handle_data(self, data):
        self._add_text(data)


def build_task_docx(task) -> bytes:
    """Erzeugt ein .docx (bytes) für eine Aufgabe - Titel + Meta + Content."""
    try:
        import docx
    except Exception as e:
        raise TaskExportError(f"python-docx konnte nicht geladen werden: {e}") from e
    import re

    version = task.current_version
    content = version.content if version else ""
    body = re.sub(r"<h2[^>]*>.*?</h2>", "", content, count=1, flags=re.DOTALL).strip()

    document = docx.Document()

    title_heading = document.add_heading(task.title, level=0)

    meta_parts = []
    if task.participant_count:
        meta_parts.append(f"{task.participant_count} Teilnehmende")
    if task.duration_minutes:
        meta_parts.append(f"{task.duration_minutes} Minuten")
    if meta_parts:
        meta_paragraph = document.add_paragraph(" · ".join(meta_parts))
        meta_paragraph.runs[0].italic = True

    parser = _TaskHtmlToDocx(document)
    parser.feed(body)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
