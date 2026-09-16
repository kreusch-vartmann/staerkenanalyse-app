"""
Unit-Tests für services/task_import.py (DOCX-/TXT-Import einzelner Aufgaben).
"""

import io

import pytest

from services.task_import import (
    TaskImportError,
    extract_task_content_from_file,
    parse_task_docx,
    parse_task_txt,
)


def _upload(filename: str, content: bytes):
    from werkzeug.datastructures import FileStorage

    return FileStorage(stream=io.BytesIO(content), filename=filename)


def _docx_bytes(paragraphs_with_styles):
    """Baut ein Test-DOCX: Liste von (text, style_name)-Tupeln.
    style_name z. B. 'Heading 1', 'Heading 2', 'List Number', 'List Bullet', None."""
    import docx

    document = docx.Document()
    for text, style in paragraphs_with_styles:
        document.add_paragraph(text, style=style)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.unit
class TestParseTaskDocx:
    def test_headings_and_lists_are_recognized(self):
        content = _docx_bytes(
            [
                ("Testaufgabe", "Heading 1"),
                ("Aufgabe", "Heading 2"),
                ("Ausgangssituation hier.", None),
                ("Rahmenbedingungen", "Heading 2"),
                ("Phase 1", "List Number"),
                ("Phase 2", "List Number"),
                ("Stifte", "List Bullet"),
            ]
        )
        html = parse_task_docx(_upload("test.docx", content))

        assert "<h2>Testaufgabe</h2>" in html
        assert "<h3>Aufgabe</h3>" in html
        assert "<h3>Rahmenbedingungen</h3>" in html
        assert "<ol>" in html and "<li>Phase 1</li>" in html
        assert "<ul>" in html and "<li>Stifte</li>" in html

    def test_bold_runs_preserved(self):
        import docx

        document = docx.Document()
        p = document.add_paragraph()
        p.add_run("Normal ")
        run = p.add_run("wichtig")
        run.bold = True
        buffer = io.BytesIO()
        document.save(buffer)

        html = parse_task_docx(_upload("bold.docx", buffer.getvalue()))
        assert "<strong>wichtig</strong>" in html

    def test_empty_docx_raises(self):
        content = _docx_bytes([])
        with pytest.raises(TaskImportError):
            parse_task_docx(_upload("empty.docx", content))

    def test_corrupt_docx_raises_clear_error(self):
        with pytest.raises(TaskImportError):
            parse_task_docx(_upload("broken.docx", b"kein gueltiges docx"))


@pytest.mark.unit
class TestParseTaskTxt:
    def test_first_short_line_becomes_heading(self):
        content = "Testaufgabe\n\nDas ist der erste Absatz.\n\nZweiter Absatz."
        html = parse_task_txt(_upload("test.txt", content.encode()))
        assert "<h2>Testaufgabe</h2>" in html
        assert "<p>Das ist der erste Absatz.</p>" in html
        assert "<p>Zweiter Absatz.</p>" in html

    def test_empty_file_raises(self):
        with pytest.raises(TaskImportError):
            parse_task_txt(_upload("empty.txt", b""))

    def test_cp1252_encoding_supported(self):
        content = "Überschrift\n\nText mit Umlauten: äöü ß".encode("cp1252")
        html = parse_task_txt(_upload("test.txt", content))
        assert "Überschrift" in html
        assert "äöü" in html

    def test_html_special_chars_are_escaped(self):
        content = "Titel\n\nText mit <script>alert(1)</script> & Ampersand.".encode()
        html = parse_task_txt(_upload("test.txt", content))
        assert "<script>" not in html
        assert "&lt;script&gt;" in html
        assert "&amp;" in html


@pytest.mark.unit
class TestExtractTaskContentFromFile:
    def test_unsupported_extension_raises(self):
        with pytest.raises(TaskImportError, match="nicht unterstützt"):
            extract_task_content_from_file(_upload("test.pdf", b"%PDF-1.4"))

    def test_dispatches_to_correct_parser_by_extension(self):
        html = extract_task_content_from_file(_upload("test.txt", b"Titel\n\nText"))
        assert "<h2>Titel</h2>" in html
