"""
Unit-Tests für den Namenslisten-Parser (services/import_service.py).

Deckt die Formate ab, die über die Import-Seite hochgeladen werden können:
TXT, CSV, XLSX, ODS, DOCX – inklusive Kopfzeilen, getrennter Namensspalten,
Metadaten-Spalten, Duplikaten und abweichender Zeichenkodierung.
"""

import io

import pytest
from werkzeug.datastructures import FileStorage

from services.import_service import (
    MAX_NAMES,
    ImportParseError,
    extract_names_from_file,
)


def _upload(filename: str, content: bytes) -> FileStorage:
    return FileStorage(stream=io.BytesIO(content), filename=filename)


def _xlsx(rows) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _ods(rows) -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    document = OpenDocumentSpreadsheet()
    table = Table(name="Teilnehmer")
    for row in rows:
        table_row = TableRow()
        for value in row:
            cell = TableCell(valuetype="string")
            cell.addElement(P(text=value))
            table_row.addElement(cell)
        table.addElement(table_row)
    document.spreadsheet.addElement(table)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx(paragraphs, table_rows=None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.unit
class TestTextAndCsv:
    def test_txt_one_name_per_line(self):
        content = "Anna Meier\nBen Schulz\n\n   Clara Roth   \n".encode()
        assert extract_names_from_file(_upload("l.txt", content)) == [
            "Anna Meier",
            "Ben Schulz",
            "Clara Roth",
        ]

    def test_csv_semicolon_with_header_and_split_columns(self):
        content = "Vorname;Nachname\nAnna;Meier\nBen;Schulz\n".encode()
        assert extract_names_from_file(_upload("l.csv", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_csv_comma_delimiter(self):
        content = "Vorname,Nachname\nAnna,Meier\nBen,Schulz\n".encode()
        assert extract_names_from_file(_upload("l.csv", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_csv_tab_delimiter(self):
        content = "Vorname\tNachname\nAnna\tMeier\n".encode()
        assert extract_names_from_file(_upload("l.csv", content)) == ["Anna Meier"]

    def test_csv_lastname_comma_firstname_is_reordered(self):
        content = 'Name\n"Meier, Anna"\n"Schulz, Ben"\n'.encode()
        assert extract_names_from_file(_upload("l.csv", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_csv_ignores_metadata_columns(self):
        content = "Nr;Name;E-Mail\n1;Anna Meier;anna@example.de\n2;Ben Schulz;ben@example.de\n".encode()
        assert extract_names_from_file(_upload("l.csv", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_duplicates_are_removed_case_insensitively(self):
        content = "Anna Meier\nanna meier\nBen Schulz\n".encode()
        assert extract_names_from_file(_upload("l.txt", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_cp1252_encoding_is_supported(self):
        content = "Jürgen Müller\nBjörn Groß\n".encode("cp1252")
        assert extract_names_from_file(_upload("l.txt", content)) == [
            "Jürgen Müller",
            "Björn Groß",
        ]

    def test_utf8_bom_is_stripped(self):
        content = "\ufeffAnna Meier\n".encode("utf-8")
        assert extract_names_from_file(_upload("l.csv", content)) == ["Anna Meier"]


@pytest.mark.unit
class TestOfficeFormats:
    def test_xlsx_with_header(self):
        content = _xlsx([["Vorname", "Nachname"], ["Anna", "Meier"], ["Ben", "Schulz"]])
        assert extract_names_from_file(_upload("l.xlsx", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_xlsx_numeric_index_column_ignored(self):
        content = _xlsx([["Nr", "Name"], [1, "Anna Meier"], [2, "Ben Schulz"]])
        assert extract_names_from_file(_upload("l.xlsx", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_ods_single_column(self):
        content = _ods([["Name"], ["Anna Meier"], ["Ben Schulz"]])
        assert extract_names_from_file(_upload("l.ods", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]

    def test_docx_paragraphs_and_table(self):
        content = _docx(
            paragraphs=["Anna Meier", ""],
            table_rows=[["Vorname", "Nachname"], ["Ben", "Schulz"]],
        )
        assert extract_names_from_file(_upload("l.docx", content)) == [
            "Anna Meier",
            "Ben Schulz",
        ]


@pytest.mark.unit
class TestErrorHandling:
    def test_unsupported_extension_raises(self):
        with pytest.raises(ImportParseError, match="nicht unterstützt"):
            extract_names_from_file(_upload("l.pdf", b"%PDF-1.4"))

    def test_empty_file_raises(self):
        with pytest.raises(ImportParseError, match="leer"):
            extract_names_from_file(_upload("l.csv", b""))

    def test_corrupt_xlsx_raises_parse_error(self):
        with pytest.raises(ImportParseError):
            extract_names_from_file(_upload("l.xlsx", b"kein gueltiges xlsx"))

    def test_only_headers_yields_empty_list(self):
        content = "Vorname;Nachname\n".encode()
        assert extract_names_from_file(_upload("l.csv", content)) == []

    def test_name_count_is_capped(self):
        rows = "\n".join(f"Person Nummer{i}" for i in range(MAX_NAMES + 50))
        names = extract_names_from_file(_upload("l.txt", rows.encode()))
        assert len(names) == MAX_NAMES
