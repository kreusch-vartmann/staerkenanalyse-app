"""
Integration-Tests für den Teilnehmer-Import über /import/names.

Schwerpunkt: Alle in der UI angebotenen Formate müssen tatsächlich
akzeptiert werden. Zuvor erlaubte die Allowlist nur Dokumentformate,
wodurch CSV- und Tabellen-Uploads an der Validierung scheiterten.
"""

import io

import pytest

from models import Group, Participant


def _xlsx_bytes(rows) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _ods_bytes(rows) -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    document = OpenDocumentSpreadsheet()
    table = Table(name="Teilnehmer")
    for row in rows:
        table_row = TableRow()
        for value in row:
            cell = TableCell(valuetype="string")
            cell.addElement(P(text=value))
            table_row.addElement(cell)
        table.addElement(table_row)
    document.spreadsheet.addElement(table)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_bytes(names) -> bytes:
    import docx

    document = docx.Document()
    for name in names:
        document.add_paragraph(name)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _import(client, group_name, filename, content):
    return client.post(
        "/import/names",
        data={
            "group_name": group_name,
            "name_file": (io.BytesIO(content), filename),
        },
        content_type="multipart/form-data",
        follow_redirects=False,
    )


@pytest.mark.integration
class TestImportNamesFormats:
    def test_txt_import(self, client, db):
        response = _import(
            client, "Gruppe TXT", "namen.txt", b"Anna Meier\nBen Schulz\n"
        )
        assert response.status_code == 302

        group = Group.query.filter_by(name="Gruppe TXT").first()
        assert group is not None
        names = [p.name for p in Participant.query.filter_by(group_id=group.id)]
        assert names == ["Anna Meier", "Ben Schulz"]

    def test_csv_import_with_header_and_columns(self, client, db):
        content = "Vorname;Nachname\nAnna;Meier\nBen;Schulz\n".encode()
        response = _import(client, "Gruppe CSV", "namen.csv", content)
        assert response.status_code == 302

        group = Group.query.filter_by(name="Gruppe CSV").first()
        assert group is not None
        names = sorted(p.name for p in Participant.query.filter_by(group_id=group.id))
        assert names == ["Anna Meier", "Ben Schulz"]

    def test_xlsx_import(self, client, db):
        content = _xlsx_bytes([["Vorname", "Nachname"], ["Anna", "Meier"]])
        response = _import(client, "Gruppe XLSX", "namen.xlsx", content)
        assert response.status_code == 302

        group = Group.query.filter_by(name="Gruppe XLSX").first()
        assert group is not None
        assert Participant.query.filter_by(group_id=group.id).count() == 1

    def test_ods_import(self, client, db):
        content = _ods_bytes([["Name"], ["Anna Meier"], ["Ben Schulz"]])
        response = _import(client, "Gruppe ODS", "namen.ods", content)
        assert response.status_code == 302

        group = Group.query.filter_by(name="Gruppe ODS").first()
        assert group is not None
        assert Participant.query.filter_by(group_id=group.id).count() == 2

    def test_docx_import(self, client, db):
        content = _docx_bytes(["Anna Meier", "Ben Schulz"])
        response = _import(client, "Gruppe DOCX", "namen.docx", content)
        assert response.status_code == 302

        group = Group.query.filter_by(name="Gruppe DOCX").first()
        assert group is not None
        assert Participant.query.filter_by(group_id=group.id).count() == 2


@pytest.mark.integration
class TestImportNamesRejections:
    def test_unsupported_format_is_rejected(self, client, db):
        response = _import(client, "Gruppe PDF", "namen.pdf", b"%PDF-1.4")
        assert response.status_code == 302
        assert Group.query.filter_by(name="Gruppe PDF").first() is None

    def test_file_without_names_is_rejected(self, client, db):
        response = _import(client, "Gruppe Leer", "namen.csv", "Vorname;Nachname\n".encode())
        assert response.status_code == 302
        assert Group.query.filter_by(name="Gruppe Leer").first() is None

    def test_missing_group_name_is_rejected(self, client, db):
        response = client.post(
            "/import/names",
            data={"group_name": "", "name_file": (io.BytesIO(b"Anna Meier\n"), "n.txt")},
            content_type="multipart/form-data",
            follow_redirects=False,
        )
        assert response.status_code == 302
        assert Participant.query.filter_by(name="Anna Meier").first() is None

    def test_observer_cannot_import(self, observer_client, db):
        """Import erfordert `import.run`, das Beobachter nicht besitzen."""
        response = _import(observer_client, "Gruppe Observer", "namen.txt", b"Anna Meier\n")
        assert response.status_code == 302
        assert Group.query.filter_by(name="Gruppe Observer").first() is None
